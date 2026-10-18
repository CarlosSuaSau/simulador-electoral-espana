"""Territorial references only: never add historical votes to national poll totals."""
import json, hashlib, unicodedata
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent
book=ROOT/'data/raw/Elecciones-Congreso.xlsx'
elections={d:{'votes':{},'valid':{},'mapping':{},'excluded':[]} for d in ['2019-04-28','2019-11-10']}
for row in load_workbook(book,read_only=True)['Congreso'].iter_rows(values_only=True):
    if not hasattr(row[0],'date'): continue
    d=str(row[0].date())
    if d not in elections or not isinstance(row[3],str): continue
    e=elections[d]; desc=unicodedata.normalize('NFC',row[3])
    if desc=='Votos válidos': e['valid']={str(i):int(row[i+3] or 0) for i in range(1,53)}
    if not desc.startswith('Votos ('): continue
    name=desc[7:-1]; upper=name.upper(); key=None
    if upper.endswith(' - VOX'): key='VOX'
    elif 'SOCIALIST' in upper: key='PSOE'
    elif 'PARTIDO POPULAR' in upper or 'PARTIT POPULAR' in upper: key='PP'
    elif 'PODEMOS' in upper or 'EN COMÚ PODEM' in upper or 'MÁS PAÍS' in upper or 'MÁS COMPROMÍS' in upper or upper.startswith('COMPROMÍS:'): key='SUMAR'
    elif 'BLOQUE NACIONALISTA GALEGO' in upper: key='BNG'
    elif 'EUSKAL HERRIA BILDU' in upper: key='BILDU'
    elif 'EAJ-PNV' in upper: key='PNV'
    elif 'ESQUERRA REPUBLICANA DE CATALUNYA' in upper: key='ERC'
    elif 'JUNTS PER CATALUNYA' in upper: key='JUNTS'
    elif d=='2019-04-28' and upper.startswith('NUEVA CANARIAS'): key='NC'
    elif d=='2019-04-28' and upper.startswith('COALICION CANARIA'): key='CC'
    if key is None:
        e['excluded'].append(name); continue
    e['mapping'][name]=key
    for i in range(1,53):
        votes=e['votes'].setdefault(str(i),{})
        votes[key]=votes.get(key,0)+int(row[i+3] or 0)
for e in elections.values():
    assert len(e['valid'])==52
    assert e['votes']['38']['VOX']>0
(ROOT/'src/data/reference-elections.json').write_text(json.dumps({'source':'https://descargas.interior.gob.es/datasets/resultados_electorales/Elecciones-Congreso.xlsx','sha256':hashlib.sha256(book.read_bytes()).hexdigest(),'note':'Shared CC–NC and Navarra Suma ballots excluded: cannot identify separate electorates. Unidas Podemos / Más País are geographical proxies for Sumar, not observed transfers.','elections':elections},ensure_ascii=False,indent=2),encoding='utf-8')
print('Generated two 52-province references; Vox Tenerife:',[e['votes']['38']['VOX'] for e in elections.values()])
